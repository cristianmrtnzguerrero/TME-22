import os
import io
import re
import zipfile
import tempfile
import datetime
import subprocess
from collections import defaultdict
import pandas as pd
import openpyxl
import streamlit as st

# ==========================================================
# 1. FUNCIÓN DE RELLENO (HOJA: Castellano 62353)
# ==========================================================
def rellenar_protocolo(wb, datos_equipo, equipos_medida, tecnico, cliente, fecha_asignada, num_ot):
    HOJA_OBJETIVO = "Castellano 62353"
    ws = wb[HOJA_OBJETIVO] if HOJA_OBJETIVO in wb.sheetnames else wb.active

    # Eliminar las demás hojas para que la conversión a PDF exporte solo esta hoja
    for sheet in wb.sheetnames:
        if sheet != ws.title:
            del wb[sheet]

    fecha_str = fecha_asignada.strftime("%d/%m/%Y")

    # --- Cabecera: Cliente, OT, Técnico, Fecha ---
    ws["D9"] = cliente
    ws["D10"] = num_ot
    ws["D11"] = tecnico
    ws["D12"] = fecha_str

    # --- Cabecera: Equipo, Marca, Modelo, Nº de Serie ---
    ws["N9"] = str(datos_equipo.get("Descripción Activo Fijo", "DESFIBRILADOR"))
    ws["N10"] = str(datos_equipo.get("Marca", ""))
    ws["N11"] = str(datos_equipo.get("Modelo", ""))
    ws["N12"] = str(datos_equipo.get("Nº de Serie", ""))

    # --- Cabecera: Inventario, Ubicación (Texto), GFH (Código) ---
    ws["X9"] = str(datos_equipo.get("Nº Activo Fijo", ""))
    ws["X10"] = str(datos_equipo.get("Descripción Ubicación Física", ""))
    ws["X11"] = str(datos_equipo.get("Ubicación Física", ""))

    # --- Equipos de medida (Filas 67 a 70) ---
    for idx, eq_m in enumerate(equipos_medida[:4]):
        fila = 67 + idx
        denominacion = str(eq_m.get("Denominación Activo", "")).strip()
        marca = str(eq_m.get("MARCA", "")).strip()
        modelo = str(eq_m.get("Modelo", "")).strip()
        n_serie = str(eq_m.get("Número de Serie del Activo", eq_m.get("Activo", ""))).strip()

        if not ws[f"A{fila}"].value:
            ws[f"A{fila}"] = denominacion

        ws[f"H{fila}"] = marca
        ws[f"M{fila}"] = modelo
        ws[f"R{fila}"] = n_serie

    return wb

# ==========================================================
# 2. CONVERSIÓN A PDF EN LA NUBE (LIBREOFFICE HEADLESS)
# ==========================================================
def exportar_a_pdf(excel_path, output_dir):
    try:
        subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", output_dir, excel_path],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE
        )
    except Exception as e:
        st.warning(f"Aviso en conversión PDF: {e}")

def limpiar_nombre_archivo(cadena):
    return re.sub(r'[\\/*?:"<>|]', "", str(cadena)).strip()

# ==========================================================
# 3. INTERFAZ WEB STREAMLIT
# ==========================================================
st.set_page_config(page_title="Generador de Protocolos", layout="wide", page_icon="⚡")
st.title("⚡ Generador de Protocolos de Revisión (Nube)")

col1, col2 = st.columns([1, 1])

with col1:
    st.subheader("1. Archivos")
    file_plantilla = st.file_uploader("1. Plantilla Excel (.xlsx)", type=["xlsx"])
    file_equipos = st.file_uploader("2. Listado de Desfibriladores (.xlsx)", type=["xlsx"])
    file_medidas = st.file_uploader("3. Listado de Comprobadores (.xlsx)", type=["xlsx"])

with col2:
    st.subheader("2. Parámetros de la Revisión")
    
    clientes_base = ["HSJ Alicante", "Hospital General de Alicante", "Hospital de San Vicente", "Otro (escribir nuevo)"]
    cliente_sel = st.selectbox("Cliente / Centro", clientes_base, index=0)
    if cliente_sel == "Otro (escribir nuevo)":
        cliente_final = st.text_input("Nombre del nuevo cliente:")
    else:
        cliente_final = cliente_sel

    tecnico = st.text_input("Técnico responsable", placeholder="Ej: Juan Pérez")
    rango_fechas = st.date_input("Rango de fechas para repartir las revisiones", value=[])

    equipos_medida_seleccionados = []
    if file_medidas:
        try:
            df_medidas = pd.read_excel(file_medidas)
            df_medidas["etiqueta"] = (
                df_medidas["Denominación Activo"].astype(str) + " | " +
                df_medidas["MARCA"].astype(str) + " " +
                df_medidas["Modelo"].astype(str) + " (S/N: " +
                df_medidas["Número de Serie del Activo"].astype(str) + ")"
            )
            
            defaults = []
            for idx, r in df_medidas.iterrows():
                den = str(r["Denominación Activo"]).lower()
                if "desfibrilador" in den or "seguridad electrica" in den:
                    defaults.append(r["etiqueta"])

            seleccion = st.multiselect(
                "Selecciona comprobadores (máx. 4):",
                options=df_medidas["etiqueta"].tolist(),
                default=defaults[:4] if defaults else None,
                max_selections=4
            )
            equipos_medida_seleccionados = df_medidas[df_medidas["etiqueta"].isin(seleccion)].to_dict(orient="records")
        except Exception as e:
            st.error(f"Error al leer comprobadores: {e}")

st.markdown("---")

# ==========================================================
# 4. GENERACIÓN MASIVA (ZIP: EXCEL + PDF)
# ==========================================================
st.subheader("3. Ejecutar y Descargar")

if st.button("🚀 Iniciar Generación", type="primary"):
    if not (file_plantilla and file_equipos and file_medidas):
        st.warning("⚠️ Debes subir los 3 archivos Excel.")
    elif not cliente_final:
        st.warning("⚠️ Debes indicar el cliente.")
    elif not tecnico:
        st.warning("⚠️ Debes escribir el nombre del técnico.")
    elif not isinstance(rango_fechas, (list, tuple)) or len(rango_fechas) < 2:
        st.warning("⚠️ Selecciona la fecha de inicio y de fin.")
    elif not equipos_medida_seleccionados:
        st.warning("⚠️ Selecciona al menos un equipo de medida.")
    else:
        fecha_inicio, fecha_fin = rango_fechas[0], rango_fechas[1]
        total_dias = (fecha_fin - fecha_inicio).days + 1
        dias_disponibles = [fecha_inicio + datetime.timedelta(days=i) for i in range(total_dias)]

        try:
            df_equipos = pd.read_excel(file_equipos)
            total_equipos = len(df_equipos)

            progreso = st.progress(0)
            status = st.empty()
            zip_buffer = io.BytesIO()

            contador_orden_dia = defaultdict(int)

            with tempfile.TemporaryDirectory() as tmpdir:
                temp_plantilla = os.path.join(tmpdir, "plantilla_base.xlsx")
                with open(temp_plantilla, "wb") as f:
                    f.write(file_plantilla.getbuffer())

                with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_out:
                    for idx, row in df_equipos.iterrows():
                        datos_eq = row.to_dict()

                        fecha_asignada = dias_disponibles[idx % len(dias_disponibles)]
                        contador_orden_dia[fecha_asignada] += 1
                        num_ot = f"{fecha_asignada.strftime('%d%m%Y')}_{contador_orden_dia[fecha_asignada]}"

                        serie = str(datos_eq.get("Nº de Serie", f"SIN_SERIE_{idx+1}")).strip()
                        nombre_base = limpiar_nombre_archivo(f"{serie}_{fecha_asignada.strftime('%d%m%Y')}")

                        status.text(f"Generando {idx+1}/{total_equipos}: {nombre_base} (OT: {num_ot})...")

                        # Rellenar Excel
                        wb = openpyxl.load_workbook(temp_plantilla)
                        wb = rellenar_protocolo(
                            wb, datos_eq, equipos_medida_seleccionados,
                            tecnico, cliente_final, fecha_asignada, num_ot
                        )

                        path_xlsx = os.path.join(tmpdir, f"{nombre_base}.xlsx")
                        wb.save(path_xlsx)
                        wb.close()

                        # Convertir a PDF en la nube
                        exportar_a_pdf(path_xlsx, tmpdir)
                        path_pdf = os.path.join(tmpdir, f"{nombre_base}.pdf")

                        # Empaquetar
                        zip_out.write(path_xlsx, arcname=f"Excel/{nombre_base}.xlsx")
                        if os.path.exists(path_pdf):
                            zip_out.write(path_pdf, arcname=f"PDF/{nombre_base}.pdf")

                        progreso.progress((idx + 1) / total_equipos)

                status.success("✅ ¡Generación completada!")
                st.download_button(
                    label="📥 Descargar todos los protocolos (ZIP)",
                    data=zip_buffer.getvalue(),
                    file_name="Protocolos_Generados.zip",
                    mime="application/zip"
                )

        except Exception as err:
            st.error(f"Se produjo un error: {err}")